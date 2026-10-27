"""Legge il materiale (PDF, slide, documenti, appunti) e chiede a Gemini di trasformarlo in post."""

import base64
import io
import json
import os
import re
import time
import unicodedata
from pathlib import Path

import requests

GEMINI_MODEL = os.environ.get("GEMINI_MODEL") or "gemini-flash-latest"
GEMINI_FALLBACK = os.environ.get("GEMINI_FALLBACK_MODEL") or "gemini-flash-lite-latest"
MAX_INLINE_BYTES = 15 * 1024 * 1024
MAX_TEXT_CHARS = 300_000

GOOGLE_DOC = "application/vnd.google-apps.document"
GOOGLE_SLIDES = "application/vnd.google-apps.presentation"
PDF = "application/pdf"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

SUPPORTED = {GOOGLE_DOC, GOOGLE_SLIDES, PDF, PPTX, DOCX}
TEXT_EXT = {".txt", ".md", ".markdown"}

CATEGORIES = [
    "Study design", "Bias & confounding", "Randomization & blinding", "Sample size & power",
    "Endpoints & outcomes", "Statistics", "Research ethics", "Informed consent", "GCP & regulation",
    "Data management", "Reporting & publication", "Critical appraisal", "Research methods",
]

PROMPT = """You create Instagram carousel posts for "{brand}", a page for researchers and health
professionals about clinical research on human subjects and research methodology.

You are given teaching material written by the page owner (slides, notes or a document; it may be in
Italian or another language). Turn it into self-contained educational posts IN ENGLISH.

STRICT ACCURACY RULES (this audience will notice mistakes):
- Use ONLY content that is in the material. You may rephrase and clarify, and you may define a
  standard methodological term in a widely accepted way, but do NOT add new facts.
- Never invent numbers, statistics, study results, guideline versions, regulations, author names or
  references. Mention a guideline, study or reference only if the material names it.
- No medical or clinical treatment advice. If the material has nothing about research or methodology,
  return an empty list.
- Skip slides that are only titles, agendas, logistics or thanks.

POST DESIGN:
- One clear concept per post; produce as many posts as the material genuinely supports (0-{max_posts}),
  each different from the others.
- title: a sharp hook, max 70 characters (e.g. a common misconception, a question, a crisp claim).
- slides: 3 to 5 slides (the carousel adds cover and a final takeaway slide: 5–7 slides in total), each with heading (max 55 characters), body (max 140 characters: few words, large type,
  plain language for professionals, concrete, no fluff) and image_prompt.
- image_prompt: in English, 15-35 words describing ONE simple illustration that a researcher would
  connect to the slide's idea at a glance. Describe a concrete scene with 2-4 recognizable elements
  (people, pills, envelopes, blindfolds, scales, magnifying glass, documents, arrows, groups of figures)
  and what they are doing; the scene must show the CONCEPT, not just the setting (never just a
  hospital bed, a lab or a generic doctor). Good examples:
  "a patient and a doctor both wearing blindfolds, seated at a table with two identical pill bottles";
  "a crowd of simple figures split by a large coin into two equal groups, each walking toward a pill bottle";
  "a magnifying glass revealing a hidden crack in a perfect-looking bridge".
  No text, letters, numbers, labeled charts, logos or recognizable real people.
- photo_query (per slide): 2-4 English keywords to search a real stock PHOTO that fits the slide
  (e.g. "physiotherapist knee rehabilitation", "researchers meeting whiteboard", "pills blister pack"),
  only for slides whose idea a real scene can show; "" when the idea is too abstract for a photo.
- highlight (per slide): the key idea of that slide, 2-6 words copied EXACTLY from its body (same
  words, same order); it is shown with a highlighter.
- source (per slide): short citation "Surname, Year" or "Surname et al., Year" ONLY if the slide's
  content comes from a study or paper cited in the material and listed in references; otherwise "".
- title_highlight: 1-4 words copied EXACTLY from the title, the most important part of the hook.
- takeaway: the single sentence to remember, max 130 characters.
- takeaway_highlight: 2-5 words copied EXACTLY from the takeaway.
- references: the papers, guidelines or books CITED IN THE MATERIAL that support this post, in APA 7
  style (Author, A. A., & Author, B. B. (Year). Title in sentence case. Journal Name, volume(issue),
  pages. https://doi.org/...). Copy authors, year, title and journal from the material; if an element
  is missing in the material, leave it out: NEVER invent authors, years, titles, journals or DOIs.
  Empty list if the material cites nothing relevant.
- hook: the FIRST line of the caption, max 120 characters (Instagram shows only about 125 characters
  before "more"). It must make a researcher stop scrolling: name a common mistake, a counter-intuitive
  consequence, or ask a sharp question that the post answers. Use only the material: no invented numbers,
  no clickbait, no hashtags, at most one emoji. Do not repeat the title word for word.
- caption: 3-6 short lines that come AFTER the hook and add context beyond the slides; end with a question
  that invites professionals to share their experience. Max 900 characters. No hashtags inside the caption.
- hashtags: 8-12 specific hashtags (e.g. #ClinicalTrials #Biostatistics #ResearchEthics).
- category: one of {categories}.
- source_hint: where in the material this comes from (e.g. "slides 12-15"), for the owner's review.

Already published posts (do not repeat these topics): {published}

File name of the material: {filename}"""

SCHEMA = {
    "type": "OBJECT",
    "properties": {
        "posts": {
            "type": "ARRAY",
            "items": {
                "type": "OBJECT",
                "properties": {
                    "title": {"type": "STRING"},
                    "category": {"type": "STRING"},
                    "slides": {
                        "type": "ARRAY",
                        "items": {
                            "type": "OBJECT",
                            "properties": {"heading": {"type": "STRING"}, "body": {"type": "STRING"},
                                           "image_prompt": {"type": "STRING"}, "photo_query": {"type": "STRING"},
                                           "highlight": {"type": "STRING"}, "source": {"type": "STRING"}},
                            "required": ["heading", "body", "image_prompt", "photo_query", "highlight", "source"],
                        },
                    },
                    "title_highlight": {"type": "STRING"},
                    "takeaway": {"type": "STRING"},
                    "takeaway_highlight": {"type": "STRING"},
                    "references": {"type": "ARRAY", "items": {"type": "STRING"}},
                    "hook": {"type": "STRING"},
                    "caption": {"type": "STRING"},
                    "hashtags": {"type": "ARRAY", "items": {"type": "STRING"}},
                    "source_hint": {"type": "STRING"},
                },
                "required": ["title", "title_highlight", "category", "slides", "takeaway",
                             "takeaway_highlight", "hook", "caption", "hashtags", "references"],
            },
        }
    },
    "required": ["posts"],
}


# ---------------------------------------------------------------- lettura materiale
def _download(drive, file_id):
    from googleapiclient.http import MediaIoBaseDownload
    buf = io.BytesIO()
    dl = MediaIoBaseDownload(buf, drive.files().get_media(fileId=file_id, supportsAllDrives=True))
    done = False
    while not done:
        _, done = dl.next_chunk()
    return buf.getvalue()


def _export(drive, file_id, mime):
    return drive.files().export(fileId=file_id, mimeType=mime).execute()


def _pdf_part(data):
    if len(data) <= MAX_INLINE_BYTES:
        return {"inline_data": {"mime_type": PDF, "data": base64.b64encode(data).decode()}}
    from pypdf import PdfReader  # PDF troppo grande: si invia solo il testo
    text = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    return {"text": text[:MAX_TEXT_CHARS]}


def _pptx_text(data):
    from pptx import Presentation
    out = []
    for n, slide in enumerate(Presentation(io.BytesIO(data)).slides, start=1):
        texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            texts.append("Speaker notes: " + slide.notes_slide.notes_text_frame.text)
        if texts:
            out.append(f"--- Slide {n} ---\n" + "\n".join(texts))
    return "\n\n".join(out)


def _docx_text(data):
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def material_text(part):
    """Testo semplice del materiale, per controllare che autori e anni delle reference ci siano davvero."""
    if "text" in part:
        return part["text"]
    try:
        from pypdf import PdfReader
        data = base64.b64decode(part["inline_data"]["data"])
        return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    except Exception:
        return ""


def is_supported(f):
    return f["mimeType"] in SUPPORTED or Path(f["name"]).suffix.lower() in TEXT_EXT | {".pdf", ".pptx", ".docx"}


def material_part(drive, f):
    """Restituisce il materiale nel formato richiesto da Gemini, o None se non leggibile."""
    mime, ext = f["mimeType"], Path(f["name"]).suffix.lower()
    if mime == GOOGLE_DOC:
        return {"text": _export(drive, f["id"], "text/plain").decode("utf-8", "ignore")[:MAX_TEXT_CHARS]}
    if mime == GOOGLE_SLIDES:
        return _pdf_part(_export(drive, f["id"], PDF))
    data = _download(drive, f["id"])
    if mime == PDF or ext == ".pdf":
        return _pdf_part(data)
    if mime == PPTX or ext == ".pptx":
        return {"text": _pptx_text(data)[:MAX_TEXT_CHARS]}
    if mime == DOCX or ext == ".docx":
        return {"text": _docx_text(data)[:MAX_TEXT_CHARS]}
    if ext in TEXT_EXT:
        return {"text": data.decode("utf-8", "ignore")[:MAX_TEXT_CHARS]}
    return None


# ---------------------------------------------------------------- Gemini
def _clip(text, n):
    text = " ".join(str(text).split())
    return text if len(text) <= n else text[: n - 1].rsplit(" ", 1)[0] + "…"


def _words(x):
    return [w for w in re.sub(r"[^\w]+", " ", str(x).lower()).split() if w]


def exact_phrase(phrase, text, max_words=6):
    """La frase da evidenziare, solo se compare parola per parola nel testo."""
    pw, tw = _words(phrase), _words(text)
    if not pw or len(pw) > max_words:
        return ""
    for i in range(len(tw) - len(pw) + 1):
        if tw[i:i + len(pw)] == pw:
            return " ".join(str(phrase).split())
    return ""


def _ref_key(ref):
    """(primo cognome, anno) di una reference APA, per verificarla nel materiale."""
    m = re.match(r"\s*([^\W\d_][\w'’\- ]*?),", ref)
    y = re.search(r"\((\d{4})[a-z]?[,)]", ref) or re.search(r"\b(19|20)\d{2}\b", ref)
    if not m or not y:
        return None
    year = y.group(1) if len(y.group(1)) == 4 else y.group(0)
    return m.group(1).strip().split()[-1].lower(), year


CROSSREF = "https://api.crossref.org/works"


def _ref_title(ref):
    """Titolo di una reference APA: il testo tra «(Anno). » e il punto successivo."""
    m = re.search(r"\(\d{4}[a-z]?\)\.\s*(.+?[.?!])(\s|$)", ref)
    return m.group(1).strip() if m else ""


def _simple(x):
    return re.sub(r"[^a-z0-9]+", " ", unicodedata.normalize("NFKD", str(x)).encode("ascii", "ignore").decode().lower()).strip()


def _title_in_material(title, material, share=0.75):
    words = [w for w in _simple(title).split() if len(w) > 3]
    if len(words) < 2:
        return False
    low = _simple(material)
    return sum(1 for w in words if w in low) / len(words) >= share


def _apa_from_crossref(item):
    """Reference APA 7 costruita dai metadati ufficiali di Crossref."""
    authors = []
    for a in (item.get("author") or [])[:20]:
        fam = a.get("family") or a.get("name") or ""
        initials = " ".join(f"{p[0]}." for p in re.split(r"[\s-]+", a.get("given") or "") if p)
        authors.append(f"{fam}, {initials}".strip(", "))
    if len(authors) > 1:
        auth = ", ".join(authors[:-1]) + ", & " + authors[-1]
    else:
        auth = authors[0] if authors else ""
    year = ((item.get("issued") or {}).get("date-parts") or [[None]])[0][0]
    title = " ".join((item.get("title") or [""])[0].split()).rstrip(".")
    journal = (item.get("container-title") or [""])[0]
    vol, issue, pages = item.get("volume"), item.get("issue"), item.get("page")
    tail = journal
    if vol:
        tail += f", {vol}" + (f"({issue})" if issue else "")
    if pages:
        tail += f", {pages.replace('-', '–')}"
    ref = f"{auth} ({year}). {title}." + (f" {tail}." if tail else "")
    if item.get("DOI"):
        ref += f" https://doi.org/{item['DOI']}"
    return ref


def crossref_reference(ref):
    """Cerca la reference su Crossref; se autore, anno e titolo coincidono restituisce la versione ufficiale."""
    key, title = _ref_key(ref), _ref_title(ref)
    if not key or not title:
        return None
    try:
        r = requests.get(CROSSREF, params={"query.bibliographic": ref[:300], "rows": 3},
                         headers={"User-Agent": "ClinicalResearchChaosBot/1.0"}, timeout=20)
        r.raise_for_status()
        items = r.json().get("message", {}).get("items", [])
    except Exception:
        return None
    from difflib import SequenceMatcher
    for it in items:
        it_title = (it.get("title") or [""])[0]
        year = ((it.get("issued") or {}).get("date-parts") or [[None]])[0][0]
        fam = _simple(((it.get("author") or [{}])[0]).get("family", ""))
        if (str(year) == key[1] and fam and _simple(key[0]).endswith(fam.split()[-1])
                and SequenceMatcher(None, _simple(title), _simple(it_title)).ratio() >= 0.85):
            return _apa_from_crossref(it)
    return None


def verified_references(refs, material):
    """Reference controllate in tre passi:
    1) primo autore e anno devono comparire nel materiale (niente reference inventate);
    2) se Crossref trova lo stesso articolo (autore, anno e titolo coincidono), si usa la sua versione ufficiale,
       completa di rivista, volume, pagine e DOI;
    3) altrimenti si tiene solo se anche il titolo si ritrova nel materiale, togliendo rivista e DOI
       quando non compaiono nel materiale (libri, linee guida, articoli non indicizzati)."""
    if not material:
        return []
    low = material.lower()
    out = []
    for r in refs or []:
        r = " ".join(str(r).split())
        key = _ref_key(r)
        if not key or key[0] not in low or key[1] not in material:
            continue
        official = crossref_reference(r)
        if official:
            ref = official
        else:
            title = _ref_title(r)
            if not _title_in_material(title, material):
                continue
            head = r[: r.find(title) + len(title)] if title in r else r
            rest = r[len(head):].strip()
            rest = re.sub(r"https?://doi\.org/\S+|doi:\s*\S+", "", rest).strip()
            journal = re.split(r",\s*\d", rest)[0].strip(" .")
            ref = head if not journal or not _title_in_material(journal, material, 0.6) else f"{head} {rest}".strip()
        if ref not in out:
            out.append(ref[:400])
        time.sleep(0.3)
    return out[:6]


def _source_ok(source, refs):
    """La citazione breve della slide deve corrispondere a una reference tenuta."""
    m = re.match(r"\s*([^\W\d_][\w'’\-]*)", source or "")
    y = re.search(r"(19|20)\d{2}", source or "")
    if not m or not y:
        return ""
    return source.strip() if any(_ref_key(r) == (m.group(1).lower(), y.group(0)) for r in refs) else ""


def _clean_post(p, material=""):
    refs = verified_references(p.get("references"), material)
    slides = []
    for s in p.get("slides", []):
        if not (s.get("heading") and s.get("body")):
            continue
        body = _clip(s.get("body", ""), 320)
        slides.append({"heading": _clip(s.get("heading", ""), 70), "body": body,
                       "image_prompt": _clip(s.get("image_prompt", ""), 300),
                       "photo_query": _clip(s.get("photo_query", ""), 60),
                       "highlight": exact_phrase(s.get("highlight"), body),
                       "source": _source_ok(s.get("source"), refs)})
    slides = slides[:7]
    if len(slides) < 2 or not p.get("title") or not p.get("takeaway"):
        return None
    tags = ["#" + t.lstrip("#").replace(" ", "") for t in p.get("hashtags", []) if t.strip("# ")]
    category = p.get("category") if p.get("category") in CATEGORIES else "Research methods"
    title, takeaway = _clip(p["title"], 90), _clip(p["takeaway"], 160)
    return {
        "title": title,
        "title_highlight": exact_phrase(p.get("title_highlight"), title, 4),
        "category": category,
        "slides": slides,
        "takeaway": takeaway,
        "takeaway_highlight": exact_phrase(p.get("takeaway_highlight"), takeaway, 5),
        "references": refs,
        "caption": with_hook(p.get("hook"), p.get("caption", ""), 1200),
        "hashtags": tags[:15],
        "source_hint": p.get("source_hint", ""),
    }


LAST_MODEL = ""       # modello che ha dato l'ultima risposta (per segnare le verifiche fatte con il ripiego)
NO_FALLBACK_TASKS = {"nuovi post da Drive", "post bando", "scelta immagini"}
THINKING = {}         # livello di ragionamento per attività (config «gemini_thinking»), es. {"verifica bandi": "low"}
USAGE = {}            # consumo di Gemini in questo giro: {(attività, modello): contatori}
_EXHAUSTED = set()    # modelli con la quota giornaliera finita: in questo giro non si riprovano


def record_usage(task, model, data=None, status="ok"):
    u = USAGE.setdefault(f"{task}|{model}", {"richieste": 0, "ingresso": 0, "uscita": 0, "ragionamento": 0, "respinte": 0})
    if status != "ok":
        u["respinte"] += 1
        return
    meta = (data or {}).get("usageMetadata") or {}
    if (data or {}).get("modelVersion"):
        u["versione"] = data["modelVersion"]          # versione reale dietro l'alias «-latest»
    u["richieste"] += 1
    u["ingresso"] += int(meta.get("promptTokenCount") or 0)
    u["uscita"] += int(meta.get("candidatesTokenCount") or 0)
    u["ragionamento"] += int(meta.get("thoughtsTokenCount") or 0)


CHECK_CAP = 34       # richieste Gemini al giorno oltre le quali le verifiche si fermano (config «gemini_daily_cap_checks»)


def quota_day():
    """Giorno della quota Gemini: si azzera a mezzanotte del Pacifico."""
    from datetime import datetime
    from zoneinfo import ZoneInfo
    return datetime.now(ZoneInfo("America/Los_Angeles")).date().isoformat()


def gemini_used_today(folder=None):
    """Richieste Gemini riuscite oggi (giorno della quota): file usage/ dei giri precedenti + questo giro."""
    folder = Path(folder or Path(__file__).resolve().parent / "usage")
    day, total = quota_day(), 0
    for f in folder.glob("*.json"):
        try:
            data = json.loads(f.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue
        if data.get("quota_day") == day:
            total += sum(c.get("richieste", 0) for k, c in data.get("consumo", {}).items() if "|gemini" in k)
    return total + sum(c.get("richieste", 0) for k, c in USAGE.items() if "|gemini" in k)


def checks_allowed():
    """Le verifiche di bandi e posizioni si fermano quando resta solo la riserva per i post del giorno."""
    return gemini_used_today() < CHECK_CAP


def save_usage(folder):
    """Scrive il consumo del giro in usage/AAAA-MM-GG-<giro>.json (un file per giro: nessun conflitto tra giri)."""
    if not USAGE:
        return None
    from datetime import datetime, timezone
    folder = Path(folder)
    folder.mkdir(exist_ok=True)
    run = os.environ.get("GITHUB_RUN_ID") or "locale"
    out = folder / f"{datetime.now(timezone.utc):%Y-%m-%d}-{run}.json"
    out.write_text(json.dumps({"workflow": os.environ.get("GITHUB_WORKFLOW", ""), "quota_day": quota_day(), "consumo": USAGE},
                              indent=1, ensure_ascii=False), encoding="utf-8")
    return out


_LOGGED_429 = set()


def retry_delay(r):
    """Secondi suggeriti da Gemini prima di riprovare (RetryInfo), oppure None."""
    import re as _re
    m = _re.search(r'"retryDelay"\s*:\s*"(\d+)', r.text)
    return int(m.group(1)) if m else None


def quota_exhausted(r):
    """429 per quota GIORNALIERA finita. Un 429 per troppe richieste al minuto porta un retryDelay: si aspetta."""
    if r.status_code != 429:
        return False
    t = r.text
    key = t[:80]
    if key not in _LOGGED_429:                    # il primo rifiuto di ogni tipo finisce nel log, per capire quale limite
        _LOGGED_429.add(key)
        import re as _re
        ids = _re.findall(r'"quotaId"\s*:\s*"([^"]+)"', t) + _re.findall(r'"quotaValue"\s*:\s*"([^"]+)"', t)
        print(f"::warning title=Gemini 429::{', '.join(ids) or t[:300]}")
    if "PerDay" in t or "per day" in t.lower():
        return True
    return retry_delay(r) is None and "PerMinute" not in t


def with_hook(hook, caption, limit):
    """Didascalia con l'hook in prima riga (la parte visibile prima di «altro» su Instagram)."""
    hook = " ".join(str(hook or "").split()).strip()
    caption = str(caption or "").strip()
    if len(hook) > 125:
        hook = hook[:124].rsplit(" ", 1)[0] + "…"
    if not hook or caption.lower().startswith(hook.lower()[:60]):
        return caption[:limit]
    return f"{hook}\n\n{caption}"[:limit]


def _parts_to_text(parts):
    """Estrae solo le parti testuali da una lista di parts Gemini (scarta inline_data/file_data)."""
    return "\n\n".join(p["text"] for p in parts if "text" in p)


def _groq_json(prompt_text, schema, temperature=0.5, task="altro"):
    """Chiama Groq e restituisce JSON decodificato. Usato come riserva finale.
    Groq ritira i modelli senza preavviso (llama-3.3-70b-versatile è sparito a ottobre 2026):
    si prova il primo modello disponibile dell'elenco GROQ_MODELS."""
    key = os.environ.get("GROQ_API_KEY")
    if not key:
        return None
    schema_str = json.dumps(schema, ensure_ascii=False)
    system = (
        "Rispondi SOLO con un oggetto JSON valido che rispetti esattamente questo schema JSON Schema: "
        + schema_str
        + "\nNessun testo prima o dopo il JSON."
    )
    for model in GROQ_MODELS:
        out = _groq_call(key, model, system, prompt_text, temperature, task)
        if out == "modello assente":
            print(f"  Groq: modello {model} non disponibile, provo il successivo")
            continue
        return out
    return None


GROQ_MODELS = [m.strip() for m in (os.environ.get("GROQ_MODELS") or "openai/gpt-oss-120b,qwen/qwen3.8-27b").split(",") if m.strip()]


def _groq_call(key, model, system, prompt_text, temperature, task):
    body = {
        "model": model,
        "messages": [{"role": "system", "content": system}, {"role": "user", "content": prompt_text}],
        "temperature": temperature,
        "response_format": {"type": "json_object"},
    }
    for attempt in range(3):
        try:
            r = requests.post(
                "https://api.groq.com/openai/v1/chat/completions",
                headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json"},
                json=body, timeout=120,
            )
            if r.status_code == 429 or r.status_code >= 500:
                time.sleep(30 * (attempt + 1))
                continue
            if r.status_code in (400, 404) and ("model_not_found" in r.text or "decommissioned" in r.text):
                return "modello assente"
            if not r.ok:
                print(f"::warning title=Groq non disponibile::{model} {r.status_code}: {r.text[:200]}")
                return None
            text = r.json()["choices"][0]["message"]["content"]
            result = json.loads(text)
            record_usage(task, f"groq-{model.split('/')[-1]}", status="ok")
            globals()["LAST_GROQ"] = f"groq-{model.split('/')[-1]}"
            return result
        except Exception as e:
            print(f"  Groq eccezione: {e}")
            return None
    return None


def gemini_json(parts, schema, temperature=0.5, task="altro", thinking=None):
    """Chiede a Gemini una risposta JSON conforme allo schema e la restituisce già decodificata.
    `thinking` ("minimal", "low", "medium", "high") limita il ragionamento interno dei modelli Gemini 3.
    Se Gemini è esaurito, le richieste di solo testo (non la scrittura dei post) passano a Groq."""
    key = os.environ.get("GEMINI_API_KEY")
    if not key:
        raise RuntimeError("GEMINI_API_KEY mancante")
    config = {"temperature": temperature, "responseMimeType": "application/json", "responseSchema": schema}
    thinking = thinking or THINKING.get(task)
    if thinking:
        config["thinkingConfig"] = {"thinkingLevel": thinking}
    body = {"contents": [{"parts": parts}], "generationConfig": config}
    last_error = None
    # quando la quota del modello principale finisce si passa a quello più leggero (quota separata)
    for model in dict.fromkeys([GEMINI_MODEL, GEMINI_FALLBACK]):
        if model in _EXHAUSTED:
            continue
        url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent"
        for attempt in range(3):
            r = requests.post(url, headers={"x-goog-api-key": key}, json=body, timeout=300)
            if r.status_code == 429 or r.status_code >= 500:
                last_error = f"{model} {r.status_code}: {r.text[:300]}"
                record_usage(task, model, status="respinta")
                if quota_exhausted(r):
                    _EXHAUSTED.add(model)     # quota giornaliera finita: inutile riprovare questo modello
                    break
                time.sleep(min(60, (retry_delay(r) or 20 * (attempt + 1)) + 2))
                continue
            if not r.ok:
                raise RuntimeError(f"Errore Gemini {r.status_code}: {r.text[:500]}")
            data = r.json()
            record_usage(task, model, data)
            globals()["LAST_MODEL"] = data.get("modelVersion") or model
            text = "".join(p.get("text", "") for p in data["candidates"][0]["content"]["parts"] if not p.get("thought"))
            return json.loads(text)
        print(f"  Gemini {model} non disponibile, provo il modello di riserva.")
    # ── Ripiego: Groq, solo per richieste di solo testo e mai per scrivere i post ──────────────
    # (PDF e immagini non passano a Groq; i testi dei post restano a Gemini). Le verifiche fatte con Groq
    # sono segnate e vengono rifatte con Gemini prima di finire in un post.
    text_only = all("text" in p for p in parts)
    if text_only and task not in NO_FALLBACK_TASKS and os.environ.get("GROQ_API_KEY"):
        print(f"  Gemini esaurito per «{task}», provo Groq…")
        result = _groq_json(_parts_to_text(parts), schema, temperature, task)
        if result is not None:
            global LAST_MODEL
            LAST_MODEL = globals().get("LAST_GROQ", "groq")
            return result
    raise RuntimeError(f"Gemini non disponibile: {last_error}")


def groq_json(prompt_text, schema, temperature=0.0, task="altro"):
    """Interfaccia pubblica per Groq: task logici veloci (es. controllo doppioni).
    Se GROQ_API_KEY non è disponibile, ricade su gemini_json con parts testuali."""
    result = _groq_json(prompt_text, schema, temperature, task)
    if result is not None:
        return result
    # fallback: usa Gemini con il solo testo
    return gemini_json([{"text": prompt_text}], schema, temperature, task)


def generate_posts(part, filename, brand, max_posts=12, published=()):
    prompt = PROMPT.format(brand=brand, max_posts=max_posts, filename=filename,
                           categories=", ".join(CATEGORIES),
                           published="; ".join(published) if published else "none")
    posts = gemini_json([{"text": prompt}, part], SCHEMA, task="nuovi post da Drive").get("posts", [])
    material = material_text(part)
    return [c for c in (_clean_post(p, material) for p in posts) if c][:max_posts]
